import io
import time
import zipfile

import pytest
from docx import Document

from app.ingest import extract
from app.ingest.extract import MAX_CHARS, DocumentError, extract_text
from tests.timing import measure_ms, percentile

RESUME_LINES = [
    "Jordan Example - Software Engineering Intern",
    "Skills: Python, FastAPI, PostgreSQL, Docker",
    "Built a retrieval system that answers questions over course notes",
]


def make_pdf(lines: list[str], pages: int = 1) -> bytes:
    """A minimal, valid PDF with real text, built by hand."""
    content = "BT /F1 12 Tf 72 720 Td 16 TL " + " ".join(
        f"({line}) Tj T*" for line in lines
    ) + " ET"
    page_ids = [3 + i for i in range(pages)]
    font_id, content_id = 3 + pages, 4 + pages
    objects = {
        1: "<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {pages} >>",
        font_id: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        content_id: f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    }
    for p in page_ids:
        objects[p] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>"
        )
    out = b"%PDF-1.4\n"
    offsets = {}
    for obj_id in sorted(objects):
        offsets[obj_id] = len(out)
        out += f"{obj_id} 0 obj\n{objects[obj_id]}\nendobj\n".encode()
    xref = len(out)
    count = max(objects) + 1
    out += f"xref\n0 {count}\n0000000000 65535 f \n".encode()
    for obj_id in range(1, count):
        out += f"{offsets[obj_id]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- happy paths --------------------------------------------------------------

def test_pdf_text_is_extracted():
    doc = extract_text(make_pdf(RESUME_LINES), filename="resume.pdf")
    assert doc.format == "pdf"
    assert doc.pages == 1
    assert "FastAPI" in doc.text
    assert "retrieval system" in doc.text


def test_docx_paragraphs_and_tables_are_extracted():
    data = make_docx(RESUME_LINES, table=[["Languages", "Python, SQL"], ["Cloud", "AWS"]])
    doc = extract_text(data, filename="resume.docx")
    assert doc.format == "docx"
    assert "PostgreSQL" in doc.text
    assert "Languages | Python, SQL" in doc.text


def test_plain_text_is_accepted():
    data = "\n".join(RESUME_LINES).encode()
    assert extract_text(data, filename="job.txt").format == "text"


def test_type_comes_from_content_not_filename():
    doc = extract_text(make_pdf(RESUME_LINES), filename="actually_a_pdf.docx")
    assert doc.format == "pdf"


def test_whitespace_is_normalized():
    data = "Python    and\tSQL\r\n\r\n\r\n\r\nDocker and more filler text to pass the minimum".encode()
    doc = extract_text(data, filename="x.txt")
    assert "Python and SQL" in doc.text
    assert "\n\n\n" not in doc.text


def test_long_text_is_truncated_and_flagged():
    data = ("word " * 10_000).encode()
    doc = extract_text(data, filename="long.txt")
    assert doc.truncated is True
    assert doc.chars == MAX_CHARS


# --- rejections -----------------------------------------------------------------

@pytest.mark.parametrize(
    "data, filename, code",
    [
        (b"", "empty.pdf", "empty"),
        (b"\x89PNG\r\n\x1a\nimage-bytes", "photo.png", "unsupported_type"),
        (b"%PDF-1.4 this is not really a pdf", "broken.pdf", "unreadable"),
        (b"\xff\xfe\x00bad", "bad.txt", "unreadable"),
    ],
)
def test_bad_files_are_rejected_with_a_code(data, filename, code):
    with pytest.raises(DocumentError) as exc_info:
        extract_text(data, filename=filename)
    assert exc_info.value.code == code


def test_oversized_file_is_rejected():
    with pytest.raises(DocumentError) as exc_info:
        extract_text(b"%PDF" + b"0" * (5 * 1024 * 1024), filename="big.pdf")
    assert exc_info.value.code == "too_large"


def test_scanned_pdf_without_text_is_rejected():
    with pytest.raises(DocumentError) as exc_info:
        extract_text(make_pdf([]), filename="scan.pdf")
    assert exc_info.value.code == "no_text"


def test_pdf_with_too_many_pages_is_rejected():
    with pytest.raises(DocumentError) as exc_info:
        extract_text(make_pdf(RESUME_LINES, pages=11), filename="long.pdf")
    assert exc_info.value.code == "too_long"


def test_zip_that_is_not_a_docx_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "not a word document")
    with pytest.raises(DocumentError) as exc_info:
        extract_text(buffer.getvalue(), filename="fake.docx")
    assert exc_info.value.code == "unreadable"


def test_zip_bomb_is_rejected_before_parsing(monkeypatch):
    monkeypatch.setattr(extract, "MAX_UNCOMPRESSED", 1024)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "0" * 10_000)  # compresses to almost nothing
    with pytest.raises(DocumentError) as exc_info:
        extract_text(buffer.getvalue(), filename="bomb.docx")
    assert exc_info.value.code == "too_large"


# --- metric ---------------------------------------------------------------------

@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_extraction_p95_under_50ms(fmt):
    data = make_pdf(RESUME_LINES) if fmt == "pdf" else make_docx(RESUME_LINES * 20)
    times = measure_ms(lambda: extract_text(data, filename=f"r.{fmt}"))
    p50, p95 = percentile(times, 0.50), percentile(times, 0.95)
    print(f"\n{fmt} extraction p50={p50:.2f}ms p95={p95:.2f}ms (n={len(times)})")
    assert p95 < 50