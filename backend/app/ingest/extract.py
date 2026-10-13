import io
import re
import zipfile
from dataclasses import dataclass
from typing import Literal

MAX_BYTES = 5 * 1024 * 1024            # 5 MB upload limit
MAX_UNCOMPRESSED = 50 * 1024 * 1024    # zip bomb guard for DOCX
MAX_PAGES = 10
MAX_CHARS = 20_000
MIN_CHARS = 50                         # less than this means no usable text


class DocumentError(ValueError):
    """A document we refuse to process. `code` is stable; `message` is user-facing."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    format: Literal["pdf", "docx", "text"]
    chars: int
    truncated: bool
    pages: int | None = None


def extract_text(data: bytes, *, filename: str = "") -> ExtractedDocument:
    if not data:
        raise DocumentError("empty", "The file is empty.")
    if len(data) > MAX_BYTES:
        raise DocumentError("too_large", "The file is larger than 5 MB.")

    pages = None
    if data.startswith(b"%PDF"):
        fmt = "pdf"
        text, pages = _read_pdf(data)
    elif data.startswith(b"PK\x03\x04"):
        fmt = "docx"
        text = _read_docx(data)
    elif filename.lower().endswith((".txt", ".md")):
        fmt = "text"
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            raise DocumentError("unreadable", "The text file is not valid UTF-8.")
    else:
        raise DocumentError(
            "unsupported_type", "Upload a PDF, a Word (.docx) file, or plain text."
        )

    text = normalize(text)
    if len(text) < MIN_CHARS:
        raise DocumentError(
            "no_text",
            "No readable text was found. If this is a scanned PDF, "
            "upload a text-based PDF or a Word file instead.",
        )

    truncated = len(text) > MAX_CHARS
    text = text[:MAX_CHARS]
    return ExtractedDocument(
        text=text, format=fmt, chars=len(text), truncated=truncated, pages=pages
    )


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [re.sub(r"[ \t\u00a0]+", " ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _read_pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("encrypted", "The PDF is password-protected.")
        page_count = len(reader.pages)
        if page_count > MAX_PAGES:
            raise DocumentError("too_long", f"The PDF has more than {MAX_PAGES} pages.")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("unreadable", "The PDF could not be read. It may be damaged.")
    return text, page_count


def _read_docx(data: bytes) -> str:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED:
                raise DocumentError("too_large", "The document expands to an unsafe size.")
        document = Document(io.BytesIO(data))
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("unreadable", "The Word file could not be read.")

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text and cell.text not in cells:  # merged cells repeat
                    cells.append(cell.text)
            parts.append(" | ".join(cells))
    return "\n".join(parts)